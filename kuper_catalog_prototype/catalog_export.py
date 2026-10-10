"""Create store-specific CSV and Excel catalogues from normalized product rows."""
from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path

from kuper_parser import HEADERS, read_existing, save


NUMERIC_COLUMNS = {"цена_руб", "белки", "жиры", "углеводы", "ккал"}


def row_identity(item):
    store = item.get("магазин_id") or str(item.get("магазин") or "").casefold()
    product = item.get("id") or item.get("ссылка") or (
        str(item.get("название") or "").casefold() + "|" + str(item.get("фасовка") or "").casefold()
    )
    return f"{store}|{product}" if product else ""


def safe_filename(value):
    value = re.sub(r'[<>:"/\\|?*]+', "_", str(value)).strip(" .")
    return value or "Неизвестный магазин"


def spreadsheet_value(header, value):
    if header not in NUMERIC_COLUMNS or value in (None, ""):
        return "" if value is None else value
    try:
        return float(str(value).replace(" ", "").replace(",", "."))
    except ValueError:
        return value


def export_xlsx(path, rows, sheet_name="Каталог"):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise RuntimeError("Для Excel установите openpyxl: python -m pip install openpyxl") from exc

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    headers = HEADERS + list(dict.fromkeys(k for row in rows for k in row if k not in HEADERS))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = re.sub(r"[\\/*?:\[\]]", "_", sheet_name)[:31] or "Каталог"
    sheet.append(headers)
    for row in sorted(rows, key=lambda x: (x.get("категория", "").casefold(), x.get("название", "").casefold())):
        sheet.append([spreadsheet_value(header, row.get(header, "")) for header in headers])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    widths = {
        "магазин": 18, "магазин_id": 12, "id": 14, "категория": 28,
        "название": 52, "бренд": 20, "фасовка": 16, "цена_руб": 13,
        "ссылка": 35, "доступность": 13, "белки": 10, "жиры": 10,
        "углеводы": 12, "ккал": 10, "кбжу_на": 14, "источник_кбжу": 20,
        "кбжу_статус": 28, "кбжу_проверено_utc": 20, "кбжу_парсер": 18, "состав": 60,
        "использовать": 14, "заметка": 30,
        "обновлено_utc": 18,
    }
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = widths.get(header, 16)
        if header in NUMERIC_COLUMNS:
            for cell in sheet.iter_cols(min_col=index, max_col=index, min_row=2):
                for item in cell:
                    item.number_format = "0.00"
    if rows:
        table = Table(displayName="CatalogTable", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)
    workbook.save(path)
    return path


def export_multi_store_xlsx(path, grouped_rows):
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError as exc:
        raise RuntimeError("Для Excel установите openpyxl: python -m pip install openpyxl") from exc

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names = set()
    for store, rows in sorted(grouped_rows.items()):
        temp = path.with_name(f".__{safe_filename(store)}.xlsx")
        export_xlsx(temp, rows, store)
        source = load_workbook(temp)
        source_sheet = source.active
        name = source_sheet.title
        base = name
        suffix = 2
        while name.casefold() in used_names:
            tail = f" {suffix}"
            name = base[:31 - len(tail)] + tail
            suffix += 1
        used_names.add(name.casefold())
        target = workbook.create_sheet(name)
        for row in source_sheet.iter_rows():
            for cell in row:
                new_cell = target[cell.coordinate]
                new_cell.value = cell.value
                if cell.has_style:
                    from copy import copy
                    new_cell.font = copy(cell.font)
                    new_cell.fill = copy(cell.fill)
                    new_cell.border = copy(cell.border)
                    new_cell.alignment = copy(cell.alignment)
                    new_cell.number_format = cell.number_format
        target.freeze_panes = source_sheet.freeze_panes
        target.sheet_view.showGridLines = False
        target.auto_filter.ref = source_sheet.auto_filter.ref
        for key, dimension in source_sheet.column_dimensions.items():
            target.column_dimensions[key].width = dimension.width
        source.close()
        temp.unlink(missing_ok=True)
    if not workbook.sheetnames:
        workbook.create_sheet("Каталог")
    workbook.save(path)
    return path


def build_catalog_files(items, output_dir, replace_stores=False):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    grouped = defaultdict(list)
    for item in items:
        grouped[item.get("магазин") or "Неизвестный магазин"].append(item)

    result = {"stores": {}, "combined": {}}
    for store, store_items in sorted(grouped.items()):
        stem = safe_filename(store)
        csv_path = output_dir / f"{stem}.csv"
        existing_store_rows = read_existing(csv_path)
        if replace_stores:
            current_keys = {row_identity(row) for row in store_items}
            existing_store_rows = [row for row in existing_store_rows if row_identity(row) in current_keys]
        save(csv_path, existing_store_rows, store_items)
        rows = read_existing(csv_path)
        xlsx_path = output_dir / f"{stem}.xlsx"
        export_xlsx(xlsx_path, rows, store)
        result["stores"][store] = {"rows": len(rows), "csv": csv_path, "xlsx": xlsx_path}

    combined_csv = output_dir / "Все магазины.csv"
    existing_combined = read_existing(combined_csv)
    if replace_stores:
        replaced_names = {str(name).casefold() for name in grouped}
        replaced_ids = {str(row.get("магазин_id") or "").casefold() for row in items if row.get("магазин_id")}
        current_keys = {row_identity(row) for row in items}
        existing_combined = [
            row for row in existing_combined
            if (
                str(row.get("магазин") or "").casefold() not in replaced_names
                and str(row.get("магазин_id") or "").casefold() not in replaced_ids
            )
            or row_identity(row) in current_keys
        ]
    save(combined_csv, existing_combined, items)
    combined_rows = read_existing(combined_csv)
    grouped_existing = defaultdict(list)
    for row in combined_rows:
        grouped_existing[row.get("магазин") or "Неизвестный магазин"].append(row)
    combined_xlsx = output_dir / "Все магазины.xlsx"
    export_multi_store_xlsx(combined_xlsx, grouped_existing)
    result["combined"] = {"rows": len(combined_rows), "csv": combined_csv, "xlsx": combined_xlsx}
    return result


def rewrite_catalog_files(rows, output_dir):
    """Rewrite every current catalogue from complete rows, including deliberate blanks."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    grouped = defaultdict(list)
    for row in rows:
        grouped[row.get("магазин") or "Неизвестный магазин"].append(dict(row))

    result = {"stores": {}, "combined": {}}
    for store, store_rows in sorted(grouped.items()):
        stem = safe_filename(store)
        csv_path = output_dir / f"{stem}.csv"
        save(csv_path, [], store_rows)
        xlsx_path = output_dir / f"{stem}.xlsx"
        export_xlsx(xlsx_path, read_existing(csv_path), store)
        result["stores"][store] = {"rows": len(store_rows), "csv": csv_path, "xlsx": xlsx_path}

    combined_csv = output_dir / "Все магазины.csv"
    save(combined_csv, [], rows)
    combined_rows = read_existing(combined_csv)
    combined_xlsx = output_dir / "Все магазины.xlsx"
    export_multi_store_xlsx(combined_xlsx, grouped)
    result["combined"] = {"rows": len(combined_rows), "csv": combined_csv, "xlsx": combined_xlsx}
    return result
